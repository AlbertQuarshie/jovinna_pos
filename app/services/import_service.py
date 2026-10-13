"""Import products from the Excel template (templates/Jovinna_Products_Template.xlsx).

Two steps, so the owner can check before anything is saved:
  rows = read_sheet(path)   -> every row with what will happen: new, update, or an error
  result = apply(rows)      -> saves the valid rows
Photos are looked up by file name in the same folder as the Excel file.
A row whose barcode already exists updates that product instead of creating a duplicate.
"""
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import select

from app import config
from app.db.database import SessionLocal
from app.db.models import Product
from app.services import image_service
from app.services import product_service as ps

# Column titles in the template (lower case, without "*" or "(KES)") -> field name
COLUMNS = {
    "name": "name", "category": "main", "subcategory": "sub", "price": "price",
    "stock qty": "stock", "low-stock alert": "reorder", "barcode": "barcode", "photo file": "photo",
}


@dataclass
class ImportRow:
    row: int                    # row number in Excel
    name: str = ""
    main: str = ""
    sub: str = ""
    price: Decimal | None = None
    stock: int = 0
    reorder: int = 5
    barcode: str = ""
    photo: Path | None = None
    action: str = "new"         # new | update | error
    problems: list[str] = field(default_factory=list)

    @property
    def ok(self) -> bool:
        return self.action != "error"


@dataclass
class ImportResult:
    created: int = 0
    updated: int = 0
    failed: list[str] = field(default_factory=list)


class SheetError(Exception):
    pass


def _header_key(text) -> str:
    t = str(text or "").replace("*", "").strip().lower()
    t = t.split("(")[0].strip()  # "price (kes)" -> "price"
    return COLUMNS.get(t, "")


def _text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)  # barcodes typed as numbers: 6941057400001.0 -> "6941057400001"
    return str(value).strip()


def _int(value, default: int, label: str, problems: list[str]) -> int:
    text = _text(value)
    if not text:
        return default
    try:
        n = int(Decimal(text.replace(",", "")))
        if n < 0:
            raise ValueError
        return n
    except (InvalidOperation, ValueError):
        problems.append(f"{label} must be a whole number")
        return default


def read_sheet(path: str | Path) -> list[ImportRow]:
    path = Path(path)
    try:
        wb = load_workbook(path, data_only=True, read_only=True)
    except Exception as e:  # noqa: BLE001
        raise SheetError(f"Couldn't open the Excel file: {e}") from e
    ws = wb["Products"] if "Products" in wb.sheetnames else wb.worksheets[0]
    rows_iter = ws.iter_rows(values_only=True)
    header = next(rows_iter, None) or []
    cols = {_header_key(h): i for i, h in enumerate(header) if _header_key(h)}
    missing = [t for t, f in (("Name", "name"), ("Category", "main"), ("Price", "price")) if f not in cols]
    if missing:
        raise SheetError(f"The sheet is missing these columns: {', '.join(missing)}. "
                           "Use the Jovinna product template.")

    with SessionLocal() as session:
        existing = set(session.scalars(select(Product.barcode)))
    mains = {m.lower(): m for m in config.MAIN_CATEGORIES}
    seen_barcodes: dict[str, int] = {}
    result = []
    for n, values in enumerate(rows_iter, start=2):
        get = lambda f: values[cols[f]] if f in cols and cols[f] < len(values) else None
        if not any(_text(v) for v in values):
            continue  # empty row
        r = ImportRow(row=n)
        p = r.problems
        r.name = _text(get("name"))
        if not r.name:
            p.append("name is missing")
        main = _text(get("main"))
        r.main = mains.get(main.lower(), "")
        if not r.main:
            p.append(f"category must be {' or '.join(config.MAIN_CATEGORIES)}" + (f" (not '{main}')" if main else ""))
        r.sub = _text(get("sub"))
        price = _text(get("price")).replace(",", "")
        try:
            r.price = Decimal(price)
            if r.price <= 0:
                p.append("price must be more than 0")
        except InvalidOperation:
            p.append("price is missing" if not price else f"price '{price}' isn't a number")
        r.stock = _int(get("stock"), 0, "stock qty", p)
        r.reorder = _int(get("reorder"), 5, "low-stock alert", p)
        r.barcode = _text(get("barcode"))
        if r.barcode:
            if r.barcode in seen_barcodes:
                p.append(f"barcode is also used on row {seen_barcodes[r.barcode]}")
            seen_barcodes.setdefault(r.barcode, n)
        photo = _text(get("photo"))
        if photo:
            r.photo = path.parent / photo
            if not r.photo.is_file():
                p.append(f"photo '{photo}' not found next to the Excel file")
        r.action = "error" if p else ("update" if r.barcode in existing else "new")
        result.append(r)
    wb.close()
    return result


def apply(rows: list[ImportRow], progress=None) -> ImportResult:
    """Save every valid row. progress(done, total) is called after each row, if given."""
    res = ImportResult()
    todo = [r for r in rows if r.ok]
    for i, r in enumerate(todo, start=1):
        try:
            with SessionLocal() as session:
                product = session.scalar(select(Product).where(Product.barcode == r.barcode)) if r.barcode else None
            image = image_service.import_image(r.photo) if r.photo else ps.KEEP
            ps.save_product(
                product.id if product else None,
                barcode=r.barcode or ps.generate_internal_barcode(), name=r.name,
                main_category=r.main, subcategory=r.sub, selling_price=r.price,
                stock_qty=r.stock, reorder_level=r.reorder, image=image,
            )
            if product:
                res.updated += 1
            else:
                res.created += 1
        except Exception as e:  # noqa: BLE001 — one bad row shouldn't stop the rest
            res.failed.append(f"Row {r.row} ({r.name}): {e}")
        if progress:
            progress(i, len(todo))
    return res
